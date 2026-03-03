"""Excel report generator."""
from __future__ import annotations

import os
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from utils.helpers import get_logger

logger = get_logger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def generate(year: int, month: int, all_metrics: dict[str, dict], output_dir: str = "reports") -> str:
    """Generate an Excel report and return the file path."""
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, f"social_report_{year}_{month:02d}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"

    # Title
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = f"Social Media Report – {date(year, month, 1).strftime('%B %Y')}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Header row
    headers = ["Platform", "Metric", "Value"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    row = 4
    for platform, metrics in all_metrics.items():
        if not metrics:
            continue
        for metric, value in metrics.items():
            ws.cell(row=row, column=1, value=platform.capitalize())
            ws.cell(row=row, column=2, value=metric.replace("_", " ").title())
            ws.cell(row=row, column=3, value=value)
            row += 1

    # Auto-size columns (skip MergedCell objects that have no column_letter)
    for col in ws.columns:
        header = next((c for c in col if hasattr(c, "column_letter")), None)
        if header is None:
            continue
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[header.column_letter].width = min(max_len + 4, 40)

    wb.save(filepath)
    logger.info("Excel report saved: %s", filepath)
    return filepath
